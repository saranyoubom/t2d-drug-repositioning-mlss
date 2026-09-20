"""
R2C3: Extracts per-well Sample/Target/Cq records from the raw Bio-Rad CFX
Maestro "Quantification Cq Results" exports, for use by
R2C3_match_well_to_replicate.py.

Inputs required (not included in this deposit -- raw instrument exports,
provided separately per journal raw-data policy; run fix_xlsx_case.py on
each first, see that script's header):
  ./cq_fixed/aug26_cq.xlsx      (Run: Aug26)
  ./cq_fixed/nov09_cq.xlsx      (Run: Nov09)
  ./cq_fixed/jun03gr1_cq.xlsx   (Run: Jun03_Gr1)
  ./cq_fixed/jun03gr2_cq.xlsx   (Run: Jun03_Gr2)
  ./cq_fixed/jun24_cq.xlsx      (Run: Jun24; this run's "Target" column
                                 combines gene and treatment, e.g. "Lef1_Ctrl")

Output: all_cq.csv (Run, Well, Gene, Treatment, Cq)
"""
import openpyxl, csv, sys, io, re
from collections import Counter
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

CQ_DIR = "./cq_fixed"
OUT = "./all_cq.csv"

GENE_FIX = {
    'GAPDH': 'Gapdh', 'ITPR1': 'Itpr1', 'RYR1': 'Ryr1', 'RYR2': 'Ryr2', 'RYR3': 'Ryr3',
}

def clean_target(t):
    t = t.strip()
    t = re.sub(r'^\d+\.\s*', '', t)  # strip leading "NN. " numbering
    t = t.replace('Ins-1', 'Ins1').replace('GLUT-2', 'Glut2').replace('GLUT2', 'Glut2')
    t = GENE_FIX.get(t.upper(), t) if t.upper() in GENE_FIX else t
    fix2 = {'ins1':'Ins1','glut2':'Glut2','tcf7':'Tcf7','tcf7l2':'Tcf7l2','lef1':'Lef1',
             'wnt2':'Wnt2','wnt2b':'Wnt2b','wnt5a':'Wnt5a','wnt5b':'Wnt5b','wnt9a':'Wnt9a',
             'hes1':'Hes1','hey1':'Hey1','glp1r':'Glp1r','gapdh':'Gapdh','kcnj11':'Kcnj11',
             'ptbp1':'Ptbp1','cacna1c':'Cacna1c','cacna1d':'Cacna1d','itpr1':'Itpr1',
             'ryr1':'Ryr1','ryr2':'Ryr2','ryr3':'Ryr3'}
    tl = t.lower()
    if tl in fix2:
        t = fix2[tl]
    return t

def clean_treat(s):
    s = (s or '').strip().upper().replace('-', '').replace(' ', '').replace('.', '')
    if 'CTRL' in s or 'CONTROL' in s or s == '1':
        return 'CTRL'
    if 'DAPT' in s:
        return 'DAPT'
    if 'DKK' in s:
        return 'DKK1'
    return None

def extract(path, run_name, target_is_combined=False):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    hdr = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    wi, ti, si, ci = hdr.index('Well'), hdr.index('Target'), hdr.index('Sample'), hdr.index('Cq')
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        t_raw = row[ti]
        if t_raw is None or t_raw == 'SYBR':
            continue
        cq = row[ci]
        well = row[wi]
        if target_is_combined:
            parts = t_raw.rsplit('_', 1)  # e.g. "Lef1_Ctrl"
            if len(parts) != 2:
                continue
            gene, treat_raw = parts
            gene = clean_target(gene)
            treat = clean_treat(treat_raw)
        else:
            gene = clean_target(t_raw)
            treat = clean_treat(row[si])
        if treat is None or cq is None:
            continue
        out.append((run_name, well, gene, treat, float(cq)))
    return out

records = []
records += extract(f"{CQ_DIR}/aug26_cq.xlsx", "Aug26")
records += extract(f"{CQ_DIR}/nov09_cq.xlsx", "Nov09")
records += extract(f"{CQ_DIR}/jun03gr1_cq.xlsx", "Jun03_Gr1")
records += extract(f"{CQ_DIR}/jun03gr2_cq.xlsx", "Jun03_Gr2")
records += extract(f"{CQ_DIR}/jun24_cq.xlsx", "Jun24", target_is_combined=True)

with open(OUT, "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["Run", "Well", "Gene", "Treatment", "Cq"])
    for r in records:
        w.writerow(r)

genes = sorted(set(r[2] for r in records))
print(f"Extracted {len(records)} Cq records across {len(genes)} genes")
print("Genes:", genes)
c = Counter((r[2], r[3]) for r in records)
for gt in sorted(c):
    print(gt, c[gt])
