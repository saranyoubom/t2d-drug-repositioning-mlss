"""
R2C3: Extracts per-cycle fluorescence for every well of the actual RT-qPCR
experimental runs from the raw Bio-Rad CFX Maestro plate exports, into one
wide CSV for downstream sliwin() amplification-efficiency estimation
(R2C3_run_sliwin_all.R).

Inputs required (not included in this deposit -- raw instrument exports,
provided separately per journal raw-data policy; run fix_xlsx_case.py on each
first, see that script's header):
  ./raw_cfx_fixed/gr_jun03.xlsx       "SYBR" sheet   (Run: Jun03_Gr1)
  ./raw_cfx_fixed/gr_jun03_gr2.xlsx   "SYBR" sheet   (Run: Jun03_Gr2)
  ./raw_cfx_fixed/gr_jun24.xlsx       "SYBR" sheet   (Run: Jun24)
  ./raw_cfx_fixed/gr_aug26.xlsx       one sheet/gene (Run: Aug26)
  ./raw_cfx_fixed/gr_nov09.xlsx       one sheet/gene (Run: Nov09)

sybr_well_gene_mapping.csv (included in this deposit) supplies the
well->gene identity for the three "SYBR"-sheet runs, whose raw export format
does not itself label which gene occupies which well.

Output: all_wells_wide.csv (Run, Well, Gene, c1..c40 per-cycle fluorescence)
"""
import openpyxl, csv, os, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

FIXED = "./raw_cfx_fixed"
WELL_GENE_MAP = "./sybr_well_gene_mapping.csv"
OUT = "./all_wells_wide.csv"

def load_sheet(path, sheet):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    data = list(ws.iter_rows(min_row=2, max_row=41, values_only=True))
    return header, data

def well_col_index(header, well_id):
    letter = well_id[0]
    num = str(int(well_id[1:]))
    target = letter + num
    return header.index(target) if target in header else None

records = []  # (run, well, gene, [rfu values cycle1..40])

# SYBR-sheet runs: gene identity from sybr_well_gene_mapping.csv
sybr_runs = {
    'Jun03_Gr1': os.path.join(FIXED, 'gr_jun03.xlsx'),
    'Jun03_Gr2': os.path.join(FIXED, 'gr_jun03_gr2.xlsx'),
    'Jun24': os.path.join(FIXED, 'gr_jun24.xlsx'),
}
well_gene_by_run = {}
with open(WELL_GENE_MAP, encoding='utf-8') as f:
    for row in csv.DictReader(f):
        well_gene_by_run.setdefault(row['Run'], []).append((row['Well'], row['Gene']))

for run, path in sybr_runs.items():
    header, data = load_sheet(path, 'SYBR')
    for well_id, gene in well_gene_by_run.get(run, []):
        col = well_col_index(header, well_id)
        if col is None:
            print(f"WARN: {run} well {well_id} not found", file=sys.stderr)
            continue
        rfu = [r[col] for r in data]
        records.append((run, well_id, gene, rfu))

# Per-gene-sheet runs: every column in the sheet is a well
per_gene_sheet_runs = {
    'Aug26': os.path.join(FIXED, 'gr_aug26.xlsx'),
    'Nov09': os.path.join(FIXED, 'gr_nov09.xlsx'),
}
GENE_NAME_FIX = {'GAPDH': 'Gapdh', 'ITPR1': 'Itpr1', 'RYR1': 'Ryr1', 'RYR2': 'Ryr2', 'RYR3': 'Ryr3'}

for run, path in per_gene_sheet_runs.items():
    wb = openpyxl.load_workbook(path, data_only=True)
    for sheet in wb.sheetnames:
        if sheet == 'Run Information':
            continue
        gene = GENE_NAME_FIX.get(sheet, sheet)
        header, data = load_sheet(path, sheet)
        wells = header[2:]
        for i, well_id in enumerate(wells):
            col = i + 2
            rfu = [r[col] for r in data]
            records.append((run, f"{sheet}:{well_id}", gene, rfu))

maxlen = max(len(r[3]) for r in records)
with open(OUT, 'w', newline='', encoding='utf-8') as f:
    w = csv.writer(f)
    w.writerow(['Run', 'Well', 'Gene'] + [f'c{i+1}' for i in range(maxlen)])
    for run, well, gene, rfu in records:
        row = [run, well, gene] + [('' if v is None else v) for v in rfu]
        w.writerow(row)

genes = sorted(set(r[2] for r in records))
print(f"Wrote {len(records)} wells across {len(genes)} genes to {OUT}")
print("Genes:", genes)
